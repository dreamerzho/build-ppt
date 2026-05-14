#!/usr/bin/env python
"""Build editable Swiss-style PPTX decks from deck_spec.json."""

from __future__ import annotations

import argparse
from copy import deepcopy
import hashlib
import json
import math
import re
import sys
import tempfile
from pathlib import Path
from typing import Any, Callable

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


SLIDE_W = 13.333
SLIDE_H = 7.5
SLIDE_WIDTH = Inches(SLIDE_W)
SLIDE_HEIGHT = Inches(SLIDE_H)

# Visual control constants. Keep renderers in absolute-canvas mode, not HTML flow mode.
SAFE_MARGIN_X = SLIDE_W * 0.08
SAFE_MARGIN_Y = SLIDE_H * 0.08
MARGIN_X = SAFE_MARGIN_X
MARGIN_TOP = SAFE_MARGIN_Y
MARGIN_BOTTOM = SAFE_MARGIN_Y
CONTENT_W = SLIDE_W - (MARGIN_X * 2)
HEADER_LABEL_POS = (SAFE_MARGIN_X, 0.4)
PAGE_NUMBER_POS = (SLIDE_W - 1.5, 0.4)
COLUMNS = 12
GUTTER = 0.15
COL_W = (SLIDE_W - (SAFE_MARGIN_X * 2) - (GUTTER * (COLUMNS - 1))) / COLUMNS

FONT_NAME = "Microsoft YaHei"
TITLE_FONT_NAME = "Microsoft YaHei"
BODY_FONT_NAME = "Microsoft YaHei Light"
TITLE_FONT_SIZE = 80
SUBTITLE_FONT_SIZE = 18
BODY_FONT_SIZE = 24
CHROME_FONT_SIZE = 10
CAPTION_FONT_SIZE = 8
HAIRLINE_WIDTH = 0.5
MAX_LINE_WIDTH = 0.75
MINIMAL_BORDER = "323232"
CARD_FILL = "F5F5F5"

LAYOUTS = {f"S{i:02d}" for i in range(1, 23)}
ABSTRACT_LAYOUTS = {
    "cover": "S01",
    "objective_list": "S02",
    "split_statement": "S03",
    "six_cells": "S04",
    "three_layers": "S05",
    "split_hero_right": "S22",
    "duo_compare": "S08",
    "timeline": "S11",
    "the_pause": "S09",
    "brief_grid": "S16",
    "image_hero": "S22",
}
EMOJI_RE = re.compile(
    "["
    "\U0001F1E6-\U0001F1FF"
    "\U0001F300-\U0001F5FF"
    "\U0001F600-\U0001F64F"
    "\U0001F680-\U0001F6FF"
    "\U0001F700-\U0001F77F"
    "\U0001F780-\U0001F7FF"
    "\U0001F800-\U0001F8FF"
    "\U0001F900-\U0001F9FF"
    "\U0001FA00-\U0001FAFF"
    "\u2600-\u27BF"
    "]+",
    flags=re.UNICODE,
)

FA_ICONS = {
    "microchip": "\uf2db",
    "laptop": "\uf108",
    "server": "\uf233",
    "code": "\uf121",
    "chart-line": "\uf201",
    "bullseye": "\uf140",
    "arrow-trend-up": "\ue097",
    "user": "\uf007",
    "users": "\uf0c0",
    "id-badge": "\uf2c1",
    "calendar": "\uf133",
    "clock": "\uf017",
    "list-check": "\uf0ae",
    "pen-nib": "\uf5ad",
    "wand": "\uf2d4",
    "layer-group": "\uf5fd",
    "gear": "\uf013",
    "shield-halved": "\uf3ed",
    "globe": "\uf0ac",
    "lightbulb": "\uf0eb",
}
FA_FONT_FAMILY = "Font Awesome 6 Free Solid"
FA_FALLBACK = "\uf0c8"

THEMES = {
    "brutalist_tech": {
        "paper": "F8FAFC",
        "ink": "0F172A",
        "grey1": "F1F5F9",
        "grey2": "334155",
        "grey3": "94A3B8",
        "accent": "CCFF00",
        "accent_on": "0F172A",
        "pattern": "E7FF66",
        "dark": "0F172A",
    },
    "swiss_classic": {
        "paper": "F5F2EB",
        "ink": "1C1C1C",
        "grey1": "EEEAE2",
        "grey2": "B4B4B4",
        "grey3": "787878",
        "accent": "E6321E",
        "accent_on": "FFFFFF",
        "pattern": "F0A096",
        "dark": "1C1C1C",
    },
    "corporate_chic": {
        "paper": "FFFFFF",
        "ink": "0A2540",
        "grey1": "F3F7FA",
        "grey2": "C8D2DC",
        "grey3": "8A96A8",
        "accent": "00D4FF",
        "accent_on": "0A2540",
        "pattern": "A8F0FF",
        "dark": "0A2540",
    },
    "lime": {
        "paper": "FAFAF8",
        "ink": "0A0A0A",
        "grey1": "F3F3F0",
        "grey2": "D8D8D4",
        "grey3": "707070",
        "accent": "B8F000",
        "accent_on": "0A0A0A",
        "pattern": "E8FF72",
        "dark": "0A0A0A",
    },
    "ikb": {
        "paper": "FAFAF8",
        "ink": "0A0A0A",
        "grey1": "F0F0EE",
        "grey2": "D4D4D2",
        "grey3": "737373",
        "accent": "002FA7",
        "accent_on": "FFFFFF",
        "pattern": "6686D8",
        "dark": "0A0A0A",
    },
    "lemon": {
        "paper": "FAFAF8",
        "ink": "0A0A0A",
        "grey1": "F0F0EE",
        "grey2": "D4D4D2",
        "grey3": "737373",
        "accent": "FFD500",
        "accent_on": "0A0A0A",
        "pattern": "FFF08A",
        "dark": "0A0A0A",
    },
    "lemon-green": {
        "paper": "FAFAF8",
        "ink": "0A0A0A",
        "grey1": "F0F0EE",
        "grey2": "D4D4D2",
        "grey3": "737373",
        "accent": "C5E803",
        "accent_on": "0A0A0A",
        "pattern": "E8FF72",
        "dark": "0A0A0A",
    },
    "safety-orange": {
        "paper": "FAFAF8",
        "ink": "0A0A0A",
        "grey1": "F0F0EE",
        "grey2": "D4D4D2",
        "grey3": "737373",
        "accent": "FF6B35",
        "accent_on": "FFFFFF",
        "pattern": "FFC3A8",
        "dark": "0A0A0A",
    },
}


def rgb(hex_value: str) -> RGBColor:
    value = hex_value.strip().lstrip("#")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


def normalize_theme_name(spec: dict[str, Any]) -> str:
    meta = spec.get("meta") if isinstance(spec.get("meta"), dict) else {}
    return field_text(spec.get("theme") or meta.get("theme"), "brutalist_tech")


def normalize_mode(spec: dict[str, Any]) -> str:
    meta = spec.get("meta") if isinstance(spec.get("meta"), dict) else {}
    mode = field_text(spec.get("mode") or meta.get("mode"), "light").lower()
    return "dark" if mode == "dark" else "light"


def raw_mode(spec: dict[str, Any]) -> str:
    meta = spec.get("meta") if isinstance(spec.get("meta"), dict) else {}
    return field_text(spec.get("mode") or meta.get("mode"), "light").lower()


def get_colors(theme_name: str, mode: str = "light") -> dict[str, str]:
    palette = THEMES[theme_name]
    colors = dict(palette)
    if mode == "dark":
        colors["paper"] = palette.get("dark", palette["ink"])
        colors["ink"] = palette.get("paper", "FFFFFF")
        colors["grey1"] = palette.get("grey2", "334155")
        colors["grey2"] = palette.get("grey3", "94A3B8")
        colors["grey3"] = palette.get("grey3", "94A3B8")
        colors["pattern"] = palette.get("grey2", "334155")
    else:
        colors["paper"] = palette.get("paper", "FFFFFF")
        colors["ink"] = palette.get("ink", "0A0A0A")
        colors["grey2"] = palette.get("grey2", "C8C8C8")
    return colors


def inch(value: float):
    return Inches(float(value))


def get_grid(start_col: int, span: int) -> tuple[float, float]:
    """Return left and width in inches for the 12-column canvas grid."""
    start_col = max(0, min(COLUMNS - 1, int(start_col)))
    span = max(1, min(COLUMNS - start_col, int(span)))
    left = SAFE_MARGIN_X + (start_col * (COL_W + GUTTER))
    width = (span * COL_W) + ((span - 1) * GUTTER)
    return left, width


def field_text(value: Any, default: str = "") -> str:
    if value is None:
        return default
    if isinstance(value, list):
        return "\n".join(str(v) for v in value if v is not None)
    return str(value)


def has_cjk(text: str) -> bool:
    return bool(re.search(r"[\u3400-\u9fff]", text))


def letterspace_label(text: str) -> str:
    text = field_text(text).upper()
    if has_cjk(text):
        return text
    return " ".join(list(text))


def contains_emoji(text: Any) -> bool:
    return bool(EMOJI_RE.search(field_text(text)))


def strip_emoji(text: str) -> str:
    return EMOJI_RE.sub("", text)


def sanitize_payload(value: Any, warnings: list[str], path: str = "$") -> Any:
    if isinstance(value, str):
        if contains_emoji(value):
            warnings.append(f"Emoji removed at {path}")
        return strip_emoji(value)
    if isinstance(value, list):
        return [sanitize_payload(item, warnings, f"{path}[{idx}]") for idx, item in enumerate(value)]
    if isinstance(value, dict):
        return {key: sanitize_payload(item, warnings, f"{path}.{key}") for key, item in value.items()}
    return value


def paginate_content(text: str, max_chars_per_slide: int = 150) -> list[str]:
    text = field_text(text).strip()
    if len(text) <= max_chars_per_slide:
        return [text] if text else []
    parts = re.split(r"(?<=[。！？.!?])\s*", text)
    pages: list[str] = []
    current = ""
    for part in parts:
        if not part:
            continue
        if len(part) > max_chars_per_slide:
            if current:
                pages.append(current.strip())
                current = ""
            pages.extend(part[i : i + max_chars_per_slide].strip() for i in range(0, len(part), max_chars_per_slide))
            continue
        if len(current) + len(part) <= max_chars_per_slide:
            current += part
        else:
            pages.append(current.strip())
            current = part
    if current:
        pages.append(current.strip())
    return [page for page in pages if page]


def normalize_abstract_slide(slide: dict[str, Any]) -> dict[str, Any]:
    normalized = deepcopy(slide)
    content = normalized.pop("content", None)
    if isinstance(content, dict):
        merged = {**content, **normalized}
        normalized = merged
    layout_type = field_text(normalized.get("layout_type"))
    if layout_type and not normalized.get("layout"):
        normalized["layout"] = ABSTRACT_LAYOUTS.get(layout_type, layout_type)
    if not normalized.get("title") and normalized.get("headline"):
        normalized["title"] = normalized["headline"]
    return normalized


def expand_slide_overflow(slide: dict[str, Any], max_chars: int = 150) -> list[dict[str, Any]]:
    body = slide.get("body")
    if not isinstance(body, str):
        return [slide]
    pages = paginate_content(body, max_chars)
    if len(pages) <= 1:
        return [slide]
    expanded = []
    for idx, page in enumerate(pages, start=1):
        clone = deepcopy(slide)
        clone["body"] = page
        clone["title"] = f"{field_text(slide.get('title'))} · {idx:02d}"
        clone["kicker"] = field_text(slide.get("kicker"), "CONTINUED")
        expanded.append(clone)
    return expanded


def normalize_spec(spec: dict[str, Any]) -> tuple[dict[str, Any], list[str]]:
    warnings: list[str] = []
    clean = sanitize_payload(deepcopy(spec), warnings)
    slides = clean.get("slides") or []
    normalized_slides: list[dict[str, Any]] = []
    for raw_slide in slides:
        if not isinstance(raw_slide, dict):
            normalized_slides.append(raw_slide)
            continue
        normalized_slides.extend(expand_slide_overflow(normalize_abstract_slide(raw_slide)))
    clean["slides"] = normalized_slides
    return clean, warnings


def item_title(item: Any, fallback: str = "") -> str:
    if isinstance(item, dict):
        return field_text(item.get("title") or item.get("label") or item.get("name"), fallback)
    return field_text(item, fallback)


def item_body(item: Any, fallback: str = "") -> str:
    if isinstance(item, dict):
        return field_text(item.get("body") or item.get("note") or item.get("description"), fallback)
    return fallback


def item_icon(item: Any, fallback: str = "microchip") -> str:
    if isinstance(item, dict):
        icon = item.get("icon")
        if icon:
            return str(icon).lower().strip().replace(" ", "-")
    return fallback


def metric_value(metric: Any) -> str:
    if isinstance(metric, dict):
        value = field_text(metric.get("value"))
        unit = field_text(metric.get("unit"))
        return f"{value}{unit}" if unit else value
    return field_text(metric)


def metric_label(metric: Any, fallback: str = "") -> str:
    if isinstance(metric, dict):
        return field_text(metric.get("label"), fallback)
    return fallback


def image_path(image: Any) -> str:
    if isinstance(image, dict):
        return field_text(image.get("path"))
    return field_text(image)


def image_caption(image: Any) -> str:
    if isinstance(image, dict):
        return field_text(image.get("caption"))
    return ""


def resolve_image(base_dir: Path, image: Any) -> Path | None:
    raw = image_path(image)
    if not raw:
        return None
    path = Path(raw)
    if not path.is_absolute():
        path = base_dir / path
    return path


def add_bg(slide, theme: dict[str, str], color_key: str = "paper") -> None:
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = rgb(theme[color_key])


def lock_text_box(shape, *, style_type: str = "body") -> None:
    tf = shape.text_frame
    tf.word_wrap = style_type != "title"
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = Inches(0)
    tf.margin_right = Inches(0)
    tf.margin_top = Inches(0)
    tf.margin_bottom = Inches(0)


def style_font_size(style_type: str, explicit_size: float | None = None) -> float:
    if style_type == "hero":
        return 88
    if style_type == "title":
        return TITLE_FONT_SIZE
    if style_type == "subtitle":
        return SUBTITLE_FONT_SIZE
    if style_type == "body":
        return BODY_FONT_SIZE
    if style_type == "chrome":
        return CHROME_FONT_SIZE
    if style_type == "caption":
        return CAPTION_FONT_SIZE
    return explicit_size if explicit_size is not None else BODY_FONT_SIZE


def font_family_for_style(style_type: str, size: float, *, bold: bool = False) -> str:
    if style_type in {"hero", "title"} or bold:
        return TITLE_FONT_NAME
    if style_type != "custom":
        return BODY_FONT_NAME
    if size >= BODY_FONT_SIZE:
        return TITLE_FONT_NAME
    return BODY_FONT_NAME


def apply_text_style(shape, *, style_type: str = "body", color: str = "0A0A0A", bold: bool | None = None) -> None:
    lock_text_box(shape, style_type=style_type)
    size = style_font_size(style_type)
    for paragraph in shape.text_frame.paragraphs:
        paragraph.alignment = PP_ALIGN.LEFT
        paragraph.space_after = Pt(0)
        paragraph.space_before = Pt(0)
        if style_type in {"hero", "title"}:
            paragraph.line_spacing = 0.85
        elif style_type == "body":
            paragraph.line_spacing = 1.4
        for run in paragraph.runs:
            if style_type == "kicker":
                run.text = letterspace_label(str(run.text))
            is_bold = (style_type in {"hero", "title"}) if bold is None else bold
            fn = font_family_for_style(style_type, size, bold=is_bold)
            run.font.name = fn
            run.font.size = Pt(size)
            run.font.bold = is_bold
            run.font.color.rgb = rgb(color)
            # Force EA typeface into the XML so PPT does not quietly fall back to system defaults on non-CN machines.
            rPr = run._r.get_or_add_rPr()
            ea = OxmlElement("a:ea")
            ea.set("typeface", fn)
            rPr.append(ea)


def apply_minimal_border(shape, *, color: str = MINIMAL_BORDER, width: float = HAIRLINE_WIDTH) -> None:
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(min(width, MAX_LINE_WIDTH))


def apply_card_style(shape) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(CARD_FILL)
    shape.line.fill.background()


def add_text(
    slide,
    text: Any,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: float = 18,
    color: str = "0A0A0A",
    bold: bool = False,
    font: str = "Arial",
    align: str = "left",
    valign: str = "top",
    uppercase: bool = False,
    style_type: str = "custom",
):
    shape = slide.shapes.add_textbox(inch(x), inch(y), inch(w), inch(h))
    tf = shape.text_frame
    tf.clear()
    lock_text_box(shape, style_type=style_type)
    tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}[valign]
    lines = field_text(text)
    if uppercase:
        lines = lines.upper()
    resolved_size = style_font_size(style_type, size)
    parts = lines.split("\n") if lines else [""]
    for idx, part in enumerate(parts):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = letterspace_label(part) if style_type == "kicker" else part
        p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[align]
        p.space_after = Pt(0)
        p.space_before = Pt(0)
        if style_type in {"hero", "title"}:
            p.line_spacing = 0.85
        elif style_type == "body":
            p.line_spacing = 1.4
        for run in p.runs:
            is_bold = bold or style_type in {"hero", "title"} or (style_type == "custom" and resolved_size >= BODY_FONT_SIZE)
            fn = font_family_for_style(style_type, resolved_size, bold=is_bold)
            run.font.name = fn
            run.font.size = Pt(resolved_size)
            run.font.bold = is_bold
            run.font.color.rgb = rgb(color)
            rPr = run._r.get_or_add_rPr()
            ea = OxmlElement("a:ea")
            ea.set("typeface", fn)
            rPr.append(ea)
    return shape


def add_rect(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    fill: str,
    line: str | None = None,
    line_width: float = 0.5,
):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, inch(x), inch(y), inch(w), inch(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if line:
        apply_minimal_border(shape, color=line, width=line_width)
    else:
        shape.line.fill.background()
    return shape


def add_rule(slide, x1: float, y1: float, x2: float, y2: float, *, color: str, width: float = 0.5):
    line = slide.shapes.add_connector(1, inch(x1), inch(y1), inch(x2), inch(y2))
    line.line.color.rgb = rgb(color)
    line.line.width = Pt(min(width, MAX_LINE_WIDTH))
    return line


def split_accent_runs(text: str, accent_terms: list[str]) -> list[tuple[str, bool]]:
    if not text or not accent_terms:
        return [(text, False)]
    runs: list[tuple[str, bool]] = []
    i = 0
    terms = sorted([t for t in accent_terms if t], key=len, reverse=True)
    while i < len(text):
        match = next((term for term in terms if text.startswith(term, i)), None)
        if match:
            runs.append((match, True))
            i += len(match)
            continue
        j = i + 1
        while j < len(text) and not any(text.startswith(term, j) for term in terms):
            j += 1
        runs.append((text[i:j], False))
        i = j
    return runs


def add_text_rich(
    slide,
    text: Any,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: float,
    color: str,
    accent_color: str,
    accent_terms: list[str] | None = None,
    font: str = "Arial",
    align: str = "left",
):
    shape = slide.shapes.add_textbox(inch(x), inch(y), inch(w), inch(h))
    tf = shape.text_frame
    tf.clear()
    is_title = size >= BODY_FONT_SIZE
    lock_text_box(shape, style_type="title" if is_title else "body")
    lines = field_text(text).split("\n") if field_text(text) else [""]
    for idx, line_text in enumerate(lines):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[align]
        p.space_after = Pt(0)
        p.space_before = Pt(0)
        if is_title:
            p.line_spacing = 0.85
        for segment, is_accent in split_accent_runs(line_text, accent_terms or []):
            run = p.add_run()
            run.text = segment
            fn = font_family_for_style("title" if is_title else "body", size, bold=is_title)
            run.font.name = fn
            run.font.size = Pt(size)
            run.font.bold = is_title
            run.font.color.rgb = rgb(accent_color if is_accent else color)
            rPr = run._r.get_or_add_rPr()
            ea = OxmlElement("a:ea")
            ea.set("typeface", fn)
            rPr.append(ea)
    return shape


def add_fa_icon(slide, icon_name: str, x: float, y: float, size: float, color: str) -> None:
    unicode_char = FA_ICONS.get(icon_name.lower(), "\uf0c8")
    box_size = size * 0.035
    shape = slide.shapes.add_textbox(inch(x), inch(y), inch(box_size), inch(box_size))
    tf = shape.text_frame
    tf.clear()
    lock_text_box(shape, style_type="custom")
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = Inches(0)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    p.space_after = Pt(0)
    p.space_before = Pt(0)
    run = p.add_run()
    run.text = unicode_char
    run.font.name = FA_FONT_FAMILY
    run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    rPr = run._r.get_or_add_rPr()
    ea = OxmlElement("a:ea")
    ea.set("typeface", FA_FONT_FAMILY)
    rPr.append(ea)


def add_micro_texture(slide, theme: dict[str, str], *, dense: bool = False) -> None:
    texture_path = build_texture_png(theme, dense=dense)
    slide.shapes.add_picture(str(texture_path), 0, 0, width=inch(SLIDE_W), height=inch(SLIDE_H))


def build_texture_png(theme: dict[str, str], *, dense: bool = False) -> Path:
    key = f"{theme.get('accent')}:{theme.get('pattern')}:{dense}:{SLIDE_W}:{SLIDE_H}"
    name = hashlib.sha1(key.encode("utf-8")).hexdigest()[:16]
    path = Path(tempfile.gettempdir()) / f"build_ppt_texture_{name}.png"
    if path.exists():
        return path
    scale = 120
    width = int(SLIDE_W * scale)
    height = int(SLIDE_H * scale)
    img = Image.new("RGBA", (width, height), (0, 0, 0, 0))
    draw = ImageDraw.Draw(img)
    pattern = theme.get("pattern", theme["grey2"]).strip().lstrip("#")
    rgba = (int(pattern[0:2], 16), int(pattern[2:4], 16), int(pattern[4:6], 16), 55)
    step = 18 if dense else 27
    for y in range(8, height - 8, step):
        for x in range(8, width - 8, step):
            if ((x // step) * 7 + (y // step) * 5) % (3 if dense else 5) != 0:
                continue
            if ((x // step) + (y // step)) % 4 == 0:
                draw.line((x - 1, y, x + 1, y), fill=rgba, width=1)
                draw.line((x, y - 1, x, y + 1), fill=rgba, width=1)
            else:
                draw.rectangle((x, y, x + 1, y + 1), fill=rgba)
    img.save(path)
    return path


def add_bottom_nav(slide, index: int, total: int, theme: dict[str, str], *, cover: bool = False) -> None:
    dot_size = 0.035
    gap = 0.07
    count = max(total, 12)
    nav_w = count * dot_size + (count - 1) * gap
    x0 = (SLIDE_W - nav_w) / 2
    y = SLIDE_H - 0.18
    inactive = theme.get("pattern", theme["grey2"]) if cover else "B7B7B3"
    active = theme["ink"] if cover else theme["accent"]
    for i in range(count):
        fill = active if i == index - 1 else inactive
        add_rect(slide, x0 + i * (dot_size + gap), y, dot_size, dot_size, fill=fill)
    hint = "← →  翻页  ·  B 静态  ·  ESC 索引"
    add_text(slide, hint, SLIDE_W - MARGIN_X - 2.35, SLIDE_H - 0.32, 2.35, 0.16, size=5.5, color=theme["ink"] if cover else theme["grey3"], align="right")


def add_pagination_nav(slide, index: int, total: int, theme: dict[str, str], *, cover: bool = False) -> None:
    dot_size = 0.035
    gap = 0.07
    count = max(total, 12)
    nav_w = count * dot_size + (count - 1) * gap
    x0 = (SLIDE_W - nav_w) / 2
    y = SLIDE_H - 0.18
    inactive = theme.get("pattern", theme["grey2"]) if cover else "B7B7B3"
    active = theme["ink"] if cover else theme["accent"]
    for i in range(count):
        fill = active if i == index - 1 else inactive
        add_rect(slide, x0 + i * (dot_size + gap), y, dot_size, dot_size, fill=fill)
    add_text(slide, "ARROW KEYS · B STATIC · ESC INDEX", SLIDE_W - MARGIN_X - 2.35, SLIDE_H - 0.32, 2.35, 0.16, size=5.5, color=theme["ink"] if cover else theme["grey3"], align="right", style_type="caption")


def add_placeholder(slide, x: float, y: float, w: float, h: float, theme: dict[str, str], label: str) -> None:
    add_rect(slide, x, y, w, h, fill=theme["grey1"], line=theme["grey2"])
    add_rule(slide, x, y, x + w, y + h, color=theme["grey2"])
    add_rule(slide, x + w, y, x, y + h, color=theme["grey2"])
    add_text(slide, label or "Missing image", x + 0.12, y + h / 2 - 0.16, w - 0.24, 0.32, size=10, color=theme["grey3"], uppercase=True)


def add_image(slide, base_dir: Path, image: Any, x: float, y: float, w: float, h: float, theme: dict[str, str], warnings: list[str]) -> None:
    path = resolve_image(base_dir, image)
    if path and path.exists():
        slide.shapes.add_picture(str(path), inch(x), inch(y), width=inch(w), height=inch(h))
        return
    raw = image_path(image) or "[empty image path]"
    warnings.append(f"Missing image: {raw}")
    add_placeholder(slide, x, y, w, h, theme, raw)


def add_masked_image(slide, base_dir: Path, image: Any, x: float, y: float, w: float, h: float, theme: dict[str, str], warnings: list[str], *, transparency: float = 0.35) -> None:
    add_image(slide, base_dir, image, x, y, w, h, theme, warnings)
    overlay = add_rect(slide, x, y, w, h, fill=theme.get("dark", theme["ink"]))
    overlay.fill.transparency = max(0, min(100, int(transparency * 100)))
    overlay.line.fill.background()


def add_rotated_meta(slide, text: str, top_pos: float, theme: dict[str, str]) -> None:
    box = add_text(slide, letterspace_label(text), 0.1, top_pos, 3.0, 0.5, size=CHROME_FONT_SIZE, color=theme["grey3"], style_type="kicker")
    box.rotation = -90.0


def add_hairline(slide, left: float, top: float, width: float, theme: dict[str, str]) -> None:
    add_rule(slide, left, top, left + width, top, color=theme["grey2"], width=HAIRLINE_WIDTH)


def inject_swipe_transition(slide) -> None:
    transition = OxmlElement("p:transition")
    transition.set("spd", "med")
    push = OxmlElement("p:push")
    push.set("dir", "l")
    transition.append(push)
    slide.element.insert(2, transition)


def add_chrome(slide, spec: dict[str, Any], slide_spec: dict[str, Any], index: int, total: int, theme: dict[str, str], *, invert: bool = False) -> None:
    color = theme["paper"] if invert else theme["grey3"]
    left = field_text(slide_spec.get("kicker") or spec.get("title") or "SWISS")
    right = f"{field_text(spec.get('author'), 'COURSE')} · {index:02d} / {total:02d}"
    add_text(slide, left, HEADER_LABEL_POS[0], HEADER_LABEL_POS[1], 6.6, 0.18, size=CHROME_FONT_SIZE, color=color, uppercase=True, style_type="chrome")
    add_text(slide, right, PAGE_NUMBER_POS[0], PAGE_NUMBER_POS[1], 1.5, 0.18, size=CHROME_FONT_SIZE, color=color, align="right", uppercase=True, style_type="chrome")


def add_notes(slide, notes: Any) -> None:
    text = field_text(notes)
    if not text:
        return
    notes_tf = slide.notes_slide.notes_text_frame
    notes_tf.clear()
    notes_tf.text = text


def add_title(slide, slide_spec: dict[str, Any], theme: dict[str, str], y: float = 0.78, h: float = 0.82, size: float = 34) -> None:
    accent_terms = slide_spec.get("accent_terms") or slide_spec.get("accent") or []
    if isinstance(accent_terms, str):
        accent_terms = [accent_terms]
    add_text_rich(
        slide,
        slide_spec.get("title", ""),
        MARGIN_X,
        y,
        CONTENT_W * 0.9,
        h,
        size=size,
        color=theme["ink"],
        accent_color=theme["accent"],
        accent_terms=accent_terms,
        font="Arial",
    )


def add_bullets(slide, items: list[Any], x: float, y: float, w: float, h: float, theme: dict[str, str], *, size: float = 14) -> None:
    shape = slide.shapes.add_textbox(inch(x), inch(y), inch(w), inch(h))
    tf = shape.text_frame
    tf.clear()
    lock_text_box(shape, style_type="body")
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        title = item_title(item)
        body = item_body(item)
        p.text = f"■  {title} - {body}" if body else f"■  {title}"
        p.level = 0
        p.font.name = BODY_FONT_NAME
        p.font.size = Pt(size)
        p.font.color.rgb = rgb(theme["accent"])
        p.space_after = Pt(10)


def render_cover(slide, spec, slide_spec, index, total, theme, base_dir, warnings):
    add_bg(slide, theme, "accent")
    add_micro_texture(slide, theme, dense=True)
    add_chrome(slide, spec, slide_spec, index, total, theme, invert=True)
    add_text(slide, field_text(slide_spec.get("eyebrow") or "AIGC · 电商物料 AI 生图 · 不止像素"), MARGIN_X, 0.72, 5.6, 0.16, size=6.8, color=theme["paper"])
    add_text(slide, slide_spec.get("kicker") or "BEYOND PIXELS · COURSE", MARGIN_X, 1.08, 5.8, 0.18, size=7.4, color=theme["ink"], uppercase=True)
    add_text_rich(
        slide,
        slide_spec.get("title"),
        MARGIN_X,
        2.28,
        8.45,
        2.38,
        size=55,
        color=theme["ink"],
        accent_color=theme["ink"],
        accent_terms=[],
        font="Arial",
    )
    add_text(slide, slide_spec.get("subtitle") or spec.get("subtitle"), MARGIN_X, 6.05, 5.8, 0.48, size=14, color=theme["ink"])
    meta = " · ".join(v for v in [field_text(slide_spec.get("author") or spec.get("author")), field_text(slide_spec.get("date") or spec.get("date"))] if v)
    add_text(slide, meta or "253班 · 2025", MARGIN_X, 6.78, 4.2, 0.16, size=6.5, color=theme["ink"], uppercase=True)


def render_timeline_kpi(slide, spec, slide_spec, index, total, theme, base_dir, warnings):
    add_bg(slide, theme)
    add_chrome(slide, spec, slide_spec, index, total, theme)
    label_x, label_w = get_grid(0, 4)
    list_x, list_w = get_grid(0, 8)
    add_text(slide, field_text(slide_spec.get("section") or "253班 · 课程目标"), label_x, 0.9, label_w, 0.18, size=6.4, color=theme["grey3"])
    add_title(slide, slide_spec, theme, y=1.18, h=1.6, size=43)
    items = slide_spec.get("items") or []
    metrics = slide_spec.get("metrics") or []
    labels = [metric_label(m) for m in metrics]
    for i, item in enumerate(items[:5]):
        y = 3.0 + i * 0.74
        label = labels[i] if i < len(labels) and labels[i] else item_title(item, f"{i + 1:02d}").upper()
        add_text(slide, f"- {i + 1:02d} / {label}", list_x, y, list_w * 0.28, 0.16, size=6.5, color=theme["accent"], uppercase=True)
        add_text(slide, item_title(item, f"Item {i + 1}"), list_x, y + 0.22, list_w * 0.7, 0.34, size=17, color=theme["ink"])
        add_text(slide, item_body(item), list_x, y + 0.58, list_w, 0.24, size=9.8, color=theme["grey3"])


def render_split_statement(slide, spec, slide_spec, index, total, theme, base_dir, warnings):
    add_bg(slide, theme)
    left_x, left_w = get_grid(0, 5)
    right_x, right_w = get_grid(6, 6)
    add_rect(slide, 0, 0, SLIDE_W * 0.48, SLIDE_H, fill=theme["ink"])
    add_text(slide, slide_spec.get("kicker") or "STATEMENT", left_x, 0.55, left_w, 0.28, size=9, color=theme["grey2"], uppercase=True)
    add_text(slide, slide_spec.get("title"), left_x, 1.2, left_w, 3.5, size=42, color=theme["paper"])
    add_chrome(slide, spec, slide_spec, index, total, theme)
    add_text(slide, slide_spec.get("body"), right_x, 1.25, right_w, 1.4, size=18, color=theme["ink"])
    add_bullets(slide, slide_spec.get("items") or [], right_x, 3.05, right_w, 2.7, theme, size=13)


def render_six_cells(slide, spec, slide_spec, index, total, theme, base_dir, warnings):
    add_bg(slide, theme)
    add_chrome(slide, spec, slide_spec, index, total, theme)
    add_title(slide, slide_spec, theme)
    items = (slide_spec.get("items") or [])[:6]
    cell_h = 1.45
    for i in range(6):
        item = items[i] if i < len(items) else {"title": f"{i + 1:02d}", "body": ""}
        col = i % 3
        row = i // 3
        x, cell_w = get_grid(col * 4, 4)
        y = 2.05 + row * (cell_h + 0.18)
        add_rect(slide, x, y, cell_w, cell_h, fill=theme["grey1"])
        pad_x, pad_y = 0.35, 0.35
        icon_name = item_icon(item)
        add_fa_icon(slide, icon_name, x + pad_x, y + 0.25, 16, theme.get("pattern", theme["grey2"]))
        add_text(slide, item_title(item), x + pad_x, y + 0.95, cell_w - pad_x*2, 0.35, size=15, color=theme["ink"], bold=True)
        add_text(slide, item_body(item), x + pad_x, y + 1.5, cell_w - pad_x*2, 0.55, size=10.5, color=theme["grey3"])


def render_three_layers(slide, spec, slide_spec, index, total, theme, base_dir, warnings):
    add_bg(slide, theme)
    add_chrome(slide, spec, slide_spec, index, total, theme)
    add_title(slide, slide_spec, theme)
    items = (slide_spec.get("items") or [])[:3]
    for i in range(3):
        item = items[i] if i < len(items) else {"title": f"Layer {i + 1}", "body": ""}
        x, card_w = get_grid(i * 4, 4)
        y = 2.0 + i * 0.25
        h = 3.55 - i * 0.25
        fill = theme["accent"] if i == 0 else theme["grey1"]
        text = theme["accent_on"] if i == 0 else theme["ink"]
        add_rect(slide, x, y, card_w, h, fill=fill)
        pad = 0.4
        icon_name = item_icon(item)
        add_fa_icon(slide, icon_name, x + pad, y + 0.2, 20, text)
        add_text(slide, item_title(item), x + pad, y + 0.75, card_w - pad*2, 0.7, size=20, color=text)
        add_text(slide, item_body(item), x + pad, y + 1.7, card_w - pad*2, 1.1, size=11, color=text)


def render_kpi_tower(slide, spec, slide_spec, index, total, theme, base_dir, warnings):
    add_bg(slide, theme)
    add_chrome(slide, spec, slide_spec, index, total, theme)
    add_title(slide, slide_spec, theme)
    add_text(slide, slide_spec.get("body"), MARGIN_X, 1.65, 4.5, 0.82, size=14, color=theme["grey3"])
    metrics = slide_spec.get("metrics") or []
    max_h = 3.3
    for i, metric in enumerate(metrics[:5]):
        x = 6.15 + i * 1.12
        raw = metric.get("max", i + 2) if isinstance(metric, dict) else i + 2
        height = max(0.55, min(max_h, float(raw) / max(5, len(metrics)) * max_h))
        add_rect(slide, x, 5.78 - height, 0.72, height, fill=theme["accent"] if i == len(metrics) - 1 else theme["ink"])
        add_text(slide, metric_value(metric), x - 0.12, 5.9, 0.96, 0.28, size=12, color=theme["ink"], align="center")
        add_text(slide, metric_label(metric), x - 0.2, 6.22, 1.12, 0.28, size=7.5, color=theme["grey3"], align="center", uppercase=True)


def render_horizontal_bar(slide, spec, slide_spec, index, total, theme, base_dir, warnings):
    add_bg(slide, theme)
    add_chrome(slide, spec, slide_spec, index, total, theme)
    add_title(slide, slide_spec, theme)
    metrics = slide_spec.get("metrics") or []
    for i, metric in enumerate(metrics[:6]):
        y = 1.85 + i * 0.62
        add_text(slide, metric_label(metric, f"Metric {i + 1}"), MARGIN_X, y, 2.4, 0.25, size=10, color=theme["ink"], uppercase=True)
        max_val = metric.get("max", 100) if isinstance(metric, dict) else 100
        value = metric.get("value", 50) if isinstance(metric, dict) else 50
        try:
            ratio = min(1, max(0.04, float(value) / float(max_val)))
        except Exception:
            ratio = 0.65
        add_rect(slide, 3.0, y, 7.2, 0.24, fill=theme["grey1"])
        add_rect(slide, 3.0, y, 7.2 * ratio, 0.24, fill=theme["accent"])
        add_text(slide, metric_value(metric), 10.45, y - 0.02, 1.3, 0.25, size=11, color=theme["ink"])


def render_duo_compare(slide, spec, slide_spec, index, total, theme, base_dir, warnings):
    add_bg(slide, theme)
    add_chrome(slide, spec, slide_spec, index, total, theme)
    add_title(slide, slide_spec, theme)
    items = slide_spec.get("items") or []
    left_x, left_w = get_grid(0, 5)
    right_x, right_w = get_grid(6, 5)
    divider_x, _ = get_grid(5, 1)
    add_rule(slide, divider_x + COL_W, 1.8, divider_x + COL_W, 6.55, color=theme["grey2"])
    for i in range(2):
        item = items[i] if i < len(items) else {"title": "Column", "body": ""}
        x, w = (left_x, left_w) if i == 0 else (right_x, right_w)
        add_text(slide, item_title(item), x, 2.0, w, 0.62, size=26, color=theme["accent"] if i == 1 else theme["ink"])
        add_text(slide, item_body(item), x, 2.92, w, 1.3, size=14, color=theme["grey3"])
        subitems = item.get("items", []) if isinstance(item, dict) else []
        add_bullets(slide, subitems, x, 4.35, w, 1.65, theme, size=11)


def render_statement(slide, spec, slide_spec, index, total, theme, base_dir, warnings):
    add_bg(slide, theme)
    add_chrome(slide, spec, slide_spec, index, total, theme)
    add_text(slide, slide_spec.get("kicker") or "STATEMENT", MARGIN_X, 1.0, 5.0, 0.28, size=9, color=theme["grey3"], uppercase=True)
    add_text(slide, slide_spec.get("title"), MARGIN_X, 1.55, 8.6, 2.3, size=46, color=theme["ink"])
    add_text(slide, slide_spec.get("subtitle") or slide_spec.get("body"), MARGIN_X, 4.35, 5.7, 0.75, size=16, color=theme["grey3"])
    for row in range(7):
        for col in range(7):
            add_rect(slide, 10.3 + col * 0.22, 1.3 + row * 0.22, 0.045, 0.045, fill=theme["accent"])


def render_split_closing(slide, spec, slide_spec, index, total, theme, base_dir, warnings):
    add_bg(slide, theme)
    add_rect(slide, SLIDE_W * 0.55, 0, SLIDE_W * 0.45, SLIDE_H, fill=theme["accent"])
    add_chrome(slide, spec, slide_spec, index, total, theme)
    add_text(slide, slide_spec.get("title"), MARGIN_X, 1.35, 5.5, 3.2, size=46, color=theme["ink"])
    add_bullets(slide, slide_spec.get("items") or [], 7.8, 1.5, 4.2, 3.9, {"ink": theme["accent_on"], "grey3": theme["accent_on"]}, size=15)


def render_horizontal_timeline(slide, spec, slide_spec, index, total, theme, base_dir, warnings):
    add_bg(slide, theme)
    add_chrome(slide, spec, slide_spec, index, total, theme)
    add_text(slide, field_text(slide_spec.get("section") or "四周路线"), MARGIN_X, 0.9, 3.0, 0.18, size=6.4, color=theme["grey3"])
    add_title(slide, slide_spec, theme, y=1.22, h=1.15, size=42)
    items = (slide_spec.get("items") or [])[:6]
    y = 4.78
    add_rule(slide, MARGIN_X + 0.7, y, SLIDE_W - MARGIN_X - 0.7, y, color=theme["grey2"], width=0.55)
    step = (CONTENT_W - 0.8) / max(1, len(items) - 1)
    for i, item in enumerate(items):
        x = MARGIN_X + 0.4 + i * step
        is_last = i == len(items) - 1
        color = theme["accent"] if is_last else theme["ink"]
        add_rect(slide, x - 0.025, y - 0.025, 0.05, 0.05, fill=color)
        add_text(slide, field_text(item.get("label") if isinstance(item, dict) else f"W{i + 1}", f"W{i + 1}"), x - 0.42, y - 0.52, 0.84, 0.18, size=8.0, color=theme["accent"] if is_last else theme["grey3"], align="center", uppercase=True)
        add_text(slide, item_title(item, f"Week {i + 1}"), x - 0.92, y - 0.28, 1.84, 0.25, size=11.5, color=color, align="center")
        body = item_body(item)
        if body:
            add_text(slide, body, x - 1.05, y + 0.24, 2.1, 0.28, size=7.5, color=theme["grey3"], align="center")


def render_manifesto(slide, spec, slide_spec, index, total, theme, base_dir, warnings):
    add_bg(slide, theme)
    add_chrome(slide, spec, slide_spec, index, total, theme)
    add_text(slide, slide_spec.get("title"), MARGIN_X, 1.05, 10.6, 2.6, size=44, color=theme["ink"])
    add_text(slide, slide_spec.get("body"), MARGIN_X, 4.05, 7.2, 0.78, size=16, color=theme["grey3"])
    add_rect(slide, 0, 5.7, SLIDE_W, 1.0, fill=theme["ink"])
    add_text(slide, slide_spec.get("subtitle") or "MANIFESTO", MARGIN_X, 5.95, CONTENT_W, 0.34, size=14, color=theme["paper"], uppercase=True)


def render_three_forces(slide, spec, slide_spec, index, total, theme, base_dir, warnings):
    add_bg(slide, theme)
    add_chrome(slide, spec, slide_spec, index, total, theme)
    items = (slide_spec.get("items") or [])[:3]
    hero_x, hero_w = get_grid(0, 4)
    card_x, card_w = get_grid(5, 7)
    add_rect(slide, hero_x, 1.1, hero_w, 5.2, fill=theme["ink"])
    add_text(slide, slide_spec.get("title"), hero_x + 0.28, 1.42, hero_w - 0.56, 2.2, size=28, color=theme["paper"])
    for i, item in enumerate(items):
        y = 1.25 + i * 1.55
        add_rect(slide, card_x, y, card_w, 1.18, fill=theme["grey1"])
        add_text(slide, item_title(item), card_x + 0.23, y + 0.2, card_w * 0.72, 0.32, size=17, color=theme["ink"])
        add_text(slide, item_body(item), card_x + 0.23, y + 0.62, card_w - 0.46, 0.28, size=10.5, color=theme["grey3"])


def render_loop(slide, spec, slide_spec, index, total, theme, base_dir, warnings):
    add_bg(slide, theme)
    add_chrome(slide, spec, slide_spec, index, total, theme)
    add_title(slide, slide_spec, theme)
    items = (slide_spec.get("items") or [])[:4]
    add_bullets(slide, items, MARGIN_X, 2.0, 4.5, 3.5, theme, size=13)
    cx, cy = 9.25, 3.8
    for i in range(4):
        angle = math.pi / 2 * i
        x = cx + math.cos(angle) * 1.25
        y = cy + math.sin(angle) * 1.25
        add_rect(slide, x - 0.28, y - 0.28, 0.56, 0.56, fill=theme["accent"] if i == 0 else theme["grey1"], line=theme["grey2"])
    add_rule(slide, cx - 1.25, cy, cx + 1.25, cy, color=theme["grey2"])
    add_rule(slide, cx, cy - 1.25, cx, cy + 1.25, color=theme["grey2"])


def render_matrix_hero(slide, spec, slide_spec, index, total, theme, base_dir, warnings):
    add_bg(slide, theme)
    add_chrome(slide, spec, slide_spec, index, total, theme)
    add_title(slide, slide_spec, theme, size=30)
    items = slide_spec.get("items") or []
    images = slide_spec.get("images") or []
    cell_w = (CONTENT_W - 0.5) / 3
    cell_h = 0.72
    for i in range(6):
        col, row = i % 3, i // 3
        x, y = MARGIN_X + col * (cell_w + 0.25), 1.85 + row * (cell_h + 0.22)
        if i < len(images):
            add_image(slide, base_dir, images[i], x, y, cell_w, cell_h, theme, warnings)
        else:
            item = items[i] if i < len(items) else {"title": f"Cell {i + 1}"}
            add_rect(slide, x, y, cell_w, cell_h, fill=theme["grey1"])
            add_text(slide, item_title(item), x + 0.12, y + 0.18, cell_w - 0.24, 0.22, size=10, color=theme["ink"], uppercase=True)
    metric = (slide_spec.get("metrics") or [{}])[0]
    add_text(slide, metric_value(metric) or "01", MARGIN_X, 4.45, 4.2, 1.05, size=54, color=theme["accent"])
    add_text(slide, metric_label(metric) or field_text(slide_spec.get("body")), MARGIN_X + 4.25, 4.82, 5.8, 0.5, size=15, color=theme["grey3"])


def render_brief(slide, spec, slide_spec, index, total, theme, base_dir, warnings):
    add_bg(slide, theme)
    add_chrome(slide, spec, slide_spec, index, total, theme)
    add_title(slide, slide_spec, theme, size=30)
    items = slide_spec.get("items") or []
    images = slide_spec.get("images") or []
    cell_w = (CONTENT_W - 0.36) / 3
    cell_h = 1.35
    for i in range(6):
        col, row = i % 3, i // 3
        x, y = MARGIN_X + col * (cell_w + 0.18), 1.9 + row * (cell_h + 0.22)
        add_rect(slide, x, y, cell_w, cell_h, fill=theme["paper"], line=theme["grey2"])
        pad = 0.25
        if i < len(images):
            add_image(slide, base_dir, images[i], x + pad, y + pad, cell_w - pad*2, 0.72, theme, warnings)
            add_text(slide, image_caption(images[i]) or item_title(items[i] if i < len(items) else ""), x + pad + 0.02, y + 0.9, cell_w - pad*2 - 0.04, 0.24, size=8.5, color=theme["grey3"], uppercase=True)
        else:
            item = items[i] if i < len(items) else {"title": f"Brief {i + 1}", "body": ""}
            add_text(slide, item_title(item), x + pad, y + 0.25, cell_w - pad*2, 0.3, size=14, color=theme["ink"])
            add_text(slide, item_body(item), x + pad, y + 0.62, cell_w - pad*2, 0.45, size=9.5, color=theme["grey3"])


def render_system(slide, spec, slide_spec, index, total, theme, base_dir, warnings):
    add_bg(slide, theme)
    add_chrome(slide, spec, slide_spec, index, total, theme)
    add_title(slide, slide_spec, theme, size=30)
    items = (slide_spec.get("items") or [])[:5]
    center_x, center_y = 6.8, 3.55
    add_rect(slide, center_x - 0.85, center_y - 0.35, 1.7, 0.7, fill=theme["accent"])
    add_text(slide, "CORE", center_x - 0.75, center_y - 0.1, 1.5, 0.22, size=11, color=theme["accent_on"], align="center", uppercase=True)
    for i, item in enumerate(items):
        angle = (2 * math.pi / max(1, len(items))) * i
        x = center_x + math.cos(angle) * 3.1
        y = center_y + math.sin(angle) * 1.65
        add_rule(slide, center_x, center_y, x, y, color=theme["grey2"])
        add_rect(slide, x - 0.75, y - 0.25, 1.5, 0.5, fill=theme["grey1"])
        add_text(slide, item_title(item), x - 0.68, y - 0.08, 1.36, 0.2, size=8.5, color=theme["ink"], align="center", uppercase=True)


def render_why_now(slide, spec, slide_spec, index, total, theme, base_dir, warnings):
    add_bg(slide, theme)
    add_chrome(slide, spec, slide_spec, index, total, theme)
    add_title(slide, slide_spec, theme)
    items = (slide_spec.get("items") or [])[:3]
    for i in range(3):
        item = items[i] if i < len(items) else {"title": f"Force {i + 1}", "body": ""}
        x = MARGIN_X + i * 3.95
        add_rule(slide, x, 2.0, x + 3.25, 2.0, color=theme["accent"] if i == 1 else theme["grey2"], width=1.2)
        add_text(slide, item_title(item), x, 2.28, 3.25, 0.5, size=19, color=theme["ink"])
        add_text(slide, item_body(item), x, 3.02, 3.1, 1.0, size=11.5, color=theme["grey3"])
    metric = (slide_spec.get("metrics") or [{}])[0]
    add_text(slide, metric_value(metric), MARGIN_X, 5.0, 4.5, 0.8, size=44, color=theme["accent"])
    add_text(slide, metric_label(metric), MARGIN_X + 4.6, 5.28, 4.5, 0.32, size=12, color=theme["grey3"], uppercase=True)


def render_four_cards(slide, spec, slide_spec, index, total, theme, base_dir, warnings):
    add_bg(slide, theme)
    add_chrome(slide, spec, slide_spec, index, total, theme)
    add_rule(slide, MARGIN_X, 0.9, SLIDE_W - MARGIN_X, 0.9, color=theme["accent"], width=2)
    add_title(slide, slide_spec, theme, y=1.08, size=30)
    items = (slide_spec.get("items") or [])[:4]
    card_w = (CONTENT_W - 0.45) / 4
    for i in range(4):
        item = items[i] if i < len(items) else {"title": f"Card {i + 1}", "body": ""}
        x = MARGIN_X + i * (card_w + 0.15)
        add_rect(slide, x, 2.25, card_w, 3.55, fill=theme["grey1"])
        add_text(slide, f"{i + 1:02d}", x + 0.14, 2.42, card_w - 0.28, 0.26, size=9, color=theme["accent"], uppercase=True)
        add_text(slide, item_title(item), x + 0.14, 2.95, card_w - 0.28, 0.7, size=18, color=theme["ink"])
        add_text(slide, item_body(item), x + 0.14, 4.0, card_w - 0.28, 0.86, size=10.5, color=theme["grey3"])


def render_ledger(slide, spec, slide_spec, index, total, theme, base_dir, warnings):
    add_bg(slide, theme)
    add_chrome(slide, spec, slide_spec, index, total, theme)
    add_title(slide, slide_spec, theme, size=30)
    metrics = slide_spec.get("metrics") or []
    for i, metric in enumerate(metrics[:4]):
        y = 1.85 + i * 1.05
        add_rule(slide, MARGIN_X, y, SLIDE_W - MARGIN_X, y, color=theme["grey2"])
        add_text(slide, metric_label(metric, f"Metric {i + 1}"), MARGIN_X, y + 0.23, 3.5, 0.24, size=10, color=theme["grey3"], uppercase=True)
        add_text(slide, metric_value(metric), 6.2, y + 0.05, 4.5, 0.62, size=34, color=theme["accent"], align="right")
        add_text(slide, field_text(metric.get("note") if isinstance(metric, dict) else ""), 10.95, y + 0.28, 1.2, 0.22, size=8, color=theme["grey3"], uppercase=True)


def render_tech_spec(slide, spec, slide_spec, index, total, theme, base_dir, warnings):
    add_bg(slide, theme)
    add_chrome(slide, spec, slide_spec, index, total, theme)
    add_title(slide, slide_spec, theme, size=34)
    metrics = slide_spec.get("metrics") or []
    for i, metric in enumerate(metrics[:3]):
        x = MARGIN_X + i * 2.7
        add_text(slide, metric_value(metric), x, 2.15, 2.3, 0.48, size=25, color=theme["accent"])
        add_text(slide, metric_label(metric), x, 2.68, 2.3, 0.24, size=8.5, color=theme["grey3"], uppercase=True)
    add_bullets(slide, slide_spec.get("items") or [], MARGIN_X, 3.45, 5.5, 2.2, theme, size=12)
    for row in range(6):
        for col in range(4):
            fill = theme["accent"] if (row + col) % 4 == 0 else theme["grey1"]
            add_rect(slide, 9.6 + col * 0.38, 3.35 + row * 0.28, 0.22, 0.16, fill=fill)


def render_image_hero(slide, spec, slide_spec, index, total, theme, base_dir, warnings):
    add_bg(slide, theme)
    add_chrome(slide, spec, slide_spec, index, total, theme)
    images = slide_spec.get("images") or []
    image = images[0] if images else {}
    if images:
        add_image(slide, base_dir, image, MARGIN_X, 1.0, CONTENT_W, 3.35, theme, warnings)
    else:
        add_placeholder(slide, MARGIN_X, 1.0, CONTENT_W, 3.35, theme, "Image slot")
    add_rect(slide, MARGIN_X + 0.25, 1.25, 4.5, 1.2, fill=theme["paper"])
    add_text(slide, slide_spec.get("title"), MARGIN_X + 0.45, 1.48, 4.05, 0.58, size=24, color=theme["ink"])
    add_text(slide, slide_spec.get("body") or image_caption(image), MARGIN_X, 4.72, 5.6, 0.46, size=13, color=theme["grey3"])
    metrics = slide_spec.get("metrics") or []
    card_w = (CONTENT_W - 0.3) / 3
    for i, metric in enumerate(metrics[:3]):
        x = MARGIN_X + i * (card_w + 0.15)
        add_text(slide, metric_value(metric), x, 5.55, card_w, 0.48, size=25, color=theme["accent"])
        add_text(slide, metric_label(metric), x, 6.02, card_w, 0.24, size=8.5, color=theme["grey3"], uppercase=True)


RENDERERS: dict[str, Callable[..., None]] = {
    "S01": render_cover,
    "S02": render_timeline_kpi,
    "S03": render_split_statement,
    "S04": render_six_cells,
    "S05": render_three_layers,
    "S06": render_kpi_tower,
    "S07": render_horizontal_bar,
    "S08": render_duo_compare,
    "S09": render_statement,
    "S10": render_split_closing,
    "S11": render_horizontal_timeline,
    "S12": render_manifesto,
    "S13": render_three_forces,
    "S14": render_loop,
    "S15": render_matrix_hero,
    "S16": render_brief,
    "S17": render_system,
    "S18": render_why_now,
    "S19": render_four_cards,
    "S20": render_ledger,
    "S21": render_tech_spec,
    "S22": render_image_hero,
}


def validate_spec(spec: dict[str, Any], base_dir: Path, *, strict_images: bool) -> tuple[list[str], list[str]]:
    errors: list[str] = []
    warnings: list[str] = []
    if not field_text(spec.get("title")):
        errors.append("Deck is missing required field: title")
    theme = normalize_theme_name(spec)
    if theme not in THEMES:
        errors.append(f"Unknown theme '{theme}'. Use one of: {', '.join(THEMES)}")
    mode = raw_mode(spec)
    if mode not in {"light", "dark"}:
        errors.append("Mode must be 'light' or 'dark'")
    slides = spec.get("slides")
    if not isinstance(slides, list) or not slides:
        errors.append("Deck must include a non-empty slides array")
        return errors, warnings
    for idx, slide in enumerate(slides, start=1):
        if not isinstance(slide, dict):
            errors.append(f"Slide {idx} must be an object")
            continue
        layout = field_text(slide.get("layout"))
        if layout not in LAYOUTS:
            errors.append(f"Slide {idx} has unknown layout '{layout}'. Use S01-S22.")
        if not field_text(slide.get("title")):
            errors.append(f"Slide {idx} ({layout or 'no layout'}) is missing required field: title")
        images = slide.get("images") or []
        if images and not isinstance(images, list):
            errors.append(f"Slide {idx} images must be an array")
            continue
        for image in images:
            path = resolve_image(base_dir, image)
            if not path or not image_path(image):
                msg = f"Slide {idx} has an empty image path"
            elif not path.exists():
                msg = f"Slide {idx} image not found: {image_path(image)}"
            else:
                continue
            if strict_images:
                errors.append(msg)
    return errors, warnings


def build_pptx(spec: dict[str, Any], base_dir: Path, out_path: Path) -> list[str]:
    theme = get_colors(normalize_theme_name(spec), normalize_mode(spec))
    warnings: list[str] = []
    prs = Presentation()
    prs.slide_width = inch(SLIDE_W)
    prs.slide_height = inch(SLIDE_H)
    blank = prs.slide_layouts[6]
    slides = spec["slides"]
    total = len(slides)
    for idx, slide_spec in enumerate(slides, start=1):
        slide = prs.slides.add_slide(blank)
        layout = field_text(slide_spec.get("layout"))
        renderer = RENDERERS[layout]
        renderer(slide, spec, slide_spec, idx, total, theme, base_dir, warnings)
        add_pagination_nav(slide, idx, total, theme, cover=layout == "S01")
        inject_swipe_transition(slide)
        add_notes(slide, slide_spec.get("notes"))
    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out_path)
    return warnings


def load_spec(path: Path) -> dict[str, Any]:
    try:
        with path.open("r", encoding="utf-8-sig") as fh:
            return json.load(fh)
    except json.JSONDecodeError as exc:
        raise SystemExit(f"Invalid JSON: {exc}") from exc


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Generate editable Swiss-style PPTX from deck_spec.json")
    parser.add_argument("spec", type=Path, help="Path to deck_spec.json")
    parser.add_argument("--out", type=Path, help="Output .pptx path")
    parser.add_argument("--validate-only", action="store_true", help="Validate the spec and image paths without writing PPTX")
    args = parser.parse_args(argv)

    spec_path = args.spec.resolve()
    spec, normalization_warnings = normalize_spec(load_spec(spec_path))
    base_dir = spec_path.parent
    errors, validation_warnings = validate_spec(spec, base_dir, strict_images=args.validate_only)
    if errors:
        for error in errors:
            print(f"ERROR: {error}", file=sys.stderr)
        return 2
    for warning in normalization_warnings + validation_warnings:
        print(f"WARNING: {warning}", file=sys.stderr)
    if args.validate_only:
        print("OK: deck spec is valid")
        return 0
    if not args.out:
        print("ERROR: --out is required unless --validate-only is used", file=sys.stderr)
        return 2
    warnings = build_pptx(spec, base_dir, args.out.resolve())
    for warning in warnings:
        print(f"WARNING: {warning}", file=sys.stderr)
    print(f"OK: wrote {args.out.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
